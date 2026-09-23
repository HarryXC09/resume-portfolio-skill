from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path('template/resume-portfolio-content-template-v2.docx')

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = tcPr.find(qn('w:shd'))
    if shd is None: shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def border(table):
    p = table._tbl.tblPr; b = p.first_child_found_in('w:tblBorders')
    if b is None: b = OxmlElement('w:tblBorders'); p.append(b)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e = b.find(qn('w:'+edge))
        if e is None: e = OxmlElement('w:'+edge); b.append(e)
        e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'6'); e.set(qn('w:color'),'D9D9D9')

def put(cell, text, bold=False, color='111827', size=9):
    cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(1); p.paragraph_format.line_spacing=1.05
    r=p.add_run(text); r.bold=bold; r.font.name='Aptos'; r.font.size=Pt(size); r.font.color.rgb=RGBColor.from_string(color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table_section(doc, title, note, rows):
    h=doc.add_paragraph(style='Heading 1'); h.add_run(title)
    p=doc.add_paragraph(note); p.paragraph_format.space_after=Pt(5)
    t=doc.add_table(rows=1, cols=4); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    widths=[1.15,1.65,3.15,1.45]
    for i,c in enumerate(t.rows[0].cells):
        c.width=Inches(widths[i]); shade(c,'243B53'); put(c,['网站区域','字段路径','可编辑内容','填写提示'][i],True,'FFFFFF',9)
    for row in rows:
        cells=t.add_row().cells
        for i,w in enumerate(widths): cells[i].width=Inches(w)
        for i,v in enumerate(row): put(cells[i],str(v),False,('0F4C5C' if i==1 else '111827'),8.5 if i!=2 else 9)
    for idx,row in enumerate(t.rows[1:],1):
        if idx%2==0:
            for c in row.cells: shade(c,'F5F8FA')
    border(t); doc.add_paragraph().paragraph_format.space_after=Pt(1)

def fields(area, mapping, tip):
    return [(area,path,f'[在此填写：{path}]',tip) for path in mapping]

D=Document(); s=D.sections[0]; s.top_margin=Inches(.55); s.bottom_margin=Inches(.55); s.left_margin=Inches(.55); s.right_margin=Inches(.55)
D.styles['Normal'].font.name='Aptos'; D.styles['Normal'].font.size=Pt(9)
for n in ('Heading 1','Heading 2'):
    D.styles[n].font.name='Aptos Display'; D.styles[n].font.color.rgb=RGBColor(0,0,0)
D.styles['Heading 1'].font.size=Pt(15); D.styles['Heading 1'].font.bold=True; D.styles['Heading 1'].paragraph_format.space_before=Pt(10); D.styles['Heading 1'].paragraph_format.space_after=Pt(4)
t=D.add_paragraph(style='Title'); r=t.add_run('人物 Portfolio 文案输入表'); r.font.name='Aptos Display'; r.font.size=Pt(24); r.font.bold=True; r.font.color.rgb=RGBColor(0,0,0)
D.add_paragraph('用于复用同一套网页简历模板。只修改“可编辑内容”列，保留“字段路径”列；Agent 会据此生成高密度初稿，再将确认后的文字同步到网页。')
D.add_paragraph('先完整填写，再人工删减。所有内容必须来自简历、作品集或本人确认的材料；中英文可先只填写一种，Agent 会保留事实边界并补齐排版所需的另一种语言。')

table_section(D,'一 个人身份与 Hero','英文名会用于首屏签名动效；职业和焦点应能在 1 秒内说明个人定位。',fields('Hero',['identity.name','identity.displayName','identity.role','identity.roleZh','identity.focus','identity.focusDetail','identity.heroTagline'],'按简历填写'))
table_section(D,'二 About 个人概览','填写 6 个能力 / 特色维度；行业不限，可以是技能、行业经验、方法论、作品方向或个人标签。',fields('About',['about.intro[0]','about.intro[1]']+[f'about.dimensions[{i}][{j}]' for i in range(6) for j in range(4)],'中文 / English 成对填写'))
# generic module selection and high-density items
mod_rows=[]
for i in range(3):
    for suffix,tip in [('type','类型：project / portfolio / award / research / publication / practice'),('labelEn','英文模块名'),('labelZh','中文模块名'),('headingEn','页面标题'),('introZh','中文导语'),('introEn','English intro')]:
        mod_rows.append((f'模块 {i+1}',f'modules[{i}].{suffix}',f'[在此填写：modules[{i}].{suffix}]',tip))
for i in range(3):
    for j in range(4):
        for suffix,tip in [('titleEn','英文条目名'),('titleZh','中文条目名'),('summaryZh','事实摘要 / 成果'),('summaryEn','English summary'),('date','时间'),('link','链接或作品编号')]:
            mod_rows.append((f'模块 {i+1} 条目 {j+1}',f'modules[{i}].items[{j}].{suffix}',f'[在此填写：modules[{i}].items[{j}].{suffix}]',tip))
table_section(D,'三 通用特色模块','不要把模板固定成学术简历。Agent 应先识别 CV 的主线，再从项目经历、个人作品集、奖项荣誉、研究、论文、客户案例、服务能力等类型中选择 1–3 个模块。初稿先填满 4 个条目，人工再删减。',mod_rows)
table_section(D,'四 教育 工作与重要经历','按时间顺序填写；机构 Logo、照片和视频编号在素材登记中处理。',fields('Journey',['journey.intro[0]','journey.intro[1]']+[f'journey.entries[{i}][{j}]' for i in range(5) for j in range(5)],'时间、机构、中文经历、英文经历'))
table_section(D,'五 精选成果与证据','该区不是固定论文区。可以填写项目、代表作品、客户案例、奖项、论文或其他最能证明能力的成果。',fields('Featured work',['featured.year','featured.title','featured.contributors','featured.meta','featured.link'],'根据个人 CV 选择最有代表性的成果'))
table_section(D,'六 联系方式','只有本人明确同意公开时才填写；手机号支持拨号链接。',fields('Contact',['contact.email','contact.wechat','contact.phone'],'可留空'))
for sec in D.sections:
    p=sec.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT; r=p.add_run('Portfolio content template'); r.font.size=Pt(8); r.font.color.rgb=RGBColor(100,116,139)
D.core_properties.title='人物 Portfolio 文案输入表'; D.core_properties.subject='通用可复用网页简历模板'; D.core_properties.author='Portfolio Template'
D.save(OUT); print(OUT)

